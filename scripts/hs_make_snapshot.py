#!/usr/bin/env python3
"""DF-17C: конвертация выгрузки TWS.by (XLSX) → JSON-снапшот для импорта.

Использование:
    python3 scripts/hs_make_snapshot.py <путь-к-tws-tnved.xlsx>

Пайплайн обновления справочника ТН ВЭД (см. docs/REFERENCES-SOURCES.md,
раздел hs):
  1) скачать свежую выгрузку: https://www.tws.by/tws/tnved/download/excel
  2) python3 scripts/hs_make_snapshot.py ~/Downloads/tws-tnved.xlsx
  3) git diff scripts/ref-data/hs/tnved-tws.json — беглая сверка версии
     и объёма (число позиций печатается скриптом)
  4) npx tsx scripts/import-hs.ts — новая редакция добавляется, старая
     закрывается (история действия сохраняется)

Скрипт валидирует выгрузку (мета-лист: дата актуальности и число кодов;
коды ровно 10 цифр; без дублей; имена непустые) и ОТКАЗЫВАЕТСЯ писать
снапшот при любом расхождении. Требуется Python 3.9+ и openpyxl
(pip install openpyxl).
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    if not src.is_file():
        print(f'ОТКАЗ: файла нет: {src}')
        return 2

    dst = Path(__file__).resolve().parent / 'ref-data' / 'hs' / 'tnved-tws.json'
    dst.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if 'ТНВЭД' not in wb.sheetnames or 'Система TWS' not in wb.sheetnames:
        print(f'ОТКАЗ: ожидались листы «ТНВЭД» и «Система TWS», есть {wb.sheetnames}')
        return 1

    # Мета-лист: «Актуальность данных дд.мм.гггг», «Всего кодов N»
    meta_date = None
    meta_count = None
    for row in wb['Система TWS'].iter_rows(values_only=True):
        a, b = row[0], row[1]
        if isinstance(a, str) and 'Актуальность' in a and b is not None:
            m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', str(b))
            if m:
                meta_date = f'{m.group(3)}-{m.group(2)}-{m.group(1)}'
        if isinstance(a, str) and 'Всего кодов' in a and b is not None:
            meta_count = int(str(b))
    if not meta_date:
        print('ОТКАЗ: в мета-листе нет «Актуальность данных» — формат выгрузки изменился?')
        return 1

    items = []
    seen = set()
    problems = []
    for row in wb['ТНВЭД'].iter_rows(min_row=2, values_only=True):
        code, name, tariff = row[0], row[1], row[2]
        if code is None:
            continue
        code = str(code).strip()
        if not re.fullmatch(r'\d{10}', code):
            problems.append((code, name))
            continue
        if code in seen:
            problems.append((code, 'ДУБЛЬ'))
            continue
        seen.add(code)
        # «родитель: 🠺 ребёнок» → «родитель / ребёнок»; схлопывание пробелов
        nm = re.sub(r'[\s:]*🠺[\s:]*', ' / ', str(name or ''))
        nm = re.sub(r'\s+', ' ', nm).strip()
        if not nm:
            problems.append((code, 'ПУСТОЕ ИМЯ'))
            continue
        tariff = str(tariff).strip() if tariff is not None else ''
        items.append({'code': code, 'name': nm, 'etdDutyText': tariff or None})

    if problems:
        print(f'ОТКАЗ: проблемы выгрузки ({len(problems)}): {problems[:10]}')
        return 1
    if meta_count is not None and meta_count != len(items):
        print(f'ОТКАЗ: мета-лист говорит {meta_count} кодов, распарсено {len(items)}')
        return 1

    snapshot = {
        'version': meta_date,
        'source': 'https://www.tws.by/tws/tnved/download/excel',
        'sourceName': 'TWS.by — открытая выгрузка справочника ТН ВЭД ЕАЭС',
        'disclaimer': 'Список кодов ТНВЭД предоставляется как есть, без каких-либо гарантий и обязательств.',
        'level': 10,
        'items': items,
    }
    dst.write_text(json.dumps(snapshot, ensure_ascii=False, indent=1), encoding='utf-8')
    size_mb = dst.stat().st_size / 1024 / 1024
    print(f'OK: редакция {meta_date}, позиций {len(items)}')
    print(f'Снапшот: {dst} ({size_mb:.2f} MB)')
    print('Далее: npx tsx scripts/import-hs.ts (новая редакция добавится, старая закроется).')
    return 0


if __name__ == '__main__':
    sys.exit(main())
